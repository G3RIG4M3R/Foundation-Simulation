#!/usr/bin/env python3
"""Independently audit final scientific evidence, PDF/PPTX equivalence and ZIP.

Optional --netlogo-home compiles the extracted model and runs the first final
XML cell once, checking the recorded seed 500001 without repeating the suite.
"""
import argparse
import csv
import hashlib
import io
import json
import math
from pathlib import Path
import statistics
import subprocess
import tempfile
import xml.etree.ElementTree as ET
import zipfile

import pymupdf
from PIL import Image
from pptx import Presentation

ROOT=Path(__file__).resolve().parents[1]
BASE=ROOT/'results/tipping_point_investigation'


def table(path):
    with path.open(newline='') as f: rows=list(csv.reader(f))
    start=next(i for i,r in enumerate(rows) if r and r[0]=='[run number]')
    header=rows[start]
    assert all(len(r)==len(header) for r in rows[start+1:] if r)
    return [dict(zip(header,r)) for r in rows[start+1:] if r]


def close(a,b): assert math.isclose(float(a),float(b),abs_tol=1e-12), (a,b)


def audit_science():
    e=ET.parse(ROOT/'experiments/tipping_confirmatory.xml').getroot().find('experiment')
    assert e.get('repetitions')=='20' and e.get('runMetricsEveryStep')=='false'
    assert e.findtext('setup')=='random-seed (500000 + behaviorspace-run-number) setup'
    assert e.find('timeLimit').get('steps')=='1000'
    assert e.findtext('exitCondition')=='ticks >= active-tick-limit'
    values={s.get('variable'):[v.get('value') for v in s] for s in e.findall('enumeratedValueSet')}
    assert {k for k,v in values.items() if len(v)>1}=={'initial-traders','trade-attractiveness'}
    assert values['initial-traders']==['8','12','16','20','24']
    assert values['trade-attractiveness']==['0.6','0.8','1']
    rows=table(BASE/'raw/confirmatory.csv')
    assert len(rows)==300 and {int(r['[run number]']) for r in rows}==set(range(1,301))
    p=json.loads((BASE/'raw/confirmatory.provenance.json').read_text())
    for key,path in [('model_sha256',ROOT/'model/foundation.nlogo'),('experiment_sha256',ROOT/'experiments/tipping_confirmatory.xml'),('csv_sha256',BASE/'raw/confirmatory.csv')]:
        assert hashlib.sha256(path.read_bytes()).hexdigest()==p[key],key
    for r in rows:
        assert r['model-valid?']=='true'
        assert float(r['ticks'])==float(r['active-tick-limit'])==1000
        assert int(r['500000 + behaviorspace-run-number'])==500000+int(r['[run number]'])
        for name,options in values.items():
            if name=='view-mode': assert r[name].strip('"')=='kingdom'
            else: assert float(r[name]) in set(map(float,options)),name
        assert 0<=float(r['control-fraction'])<=1
        close(float(r['controlled-planets'])/30,r['control-fraction'])
        assert 0<=float(r['controlled-kingdoms'])<=4
    with (BASE/'processed/confirmatory_cell_summary.csv').open() as f: summary=list(csv.DictReader(f))
    assert len(summary)==15
    for s in summary:
        samples=[r for r in rows if float(r['initial-traders'])==float(s['initial_traders']) and float(r['trade-attractiveness'])==float(s['trade_attractiveness'])]
        assert len(samples)==20
        for metric,prefix in [('control-fraction','control'),('controlled-kingdoms','kingdom')]:
            x=[float(r[metric]) for r in samples];mean=statistics.mean(x);sd=statistics.stdev(x);h=2.093024054408263*sd/math.sqrt(20)
            for suffix,val in [('mean',mean),('sd',sd),('ci_low',mean-h),('ci_high',mean+h)]:close(s[prefix+'_'+suffix],val)
        n=sum(float(r['control-fraction'])>=.5 for r in samples);close(s['widespread_count'],n)
        z=1.959963984540054;p0=n/20;d=1+z*z/20
        center=(p0+z*z/40)/d;half=z*math.sqrt(p0*(1-p0)/20+z*z/1600)/d
        close(s['widespread_wilson_low'],center-half);close(s['widespread_wilson_high'],center+half)
    with (BASE/'processed/duration_summary.csv').open() as f:
        for s in csv.DictReader(f):
            r=table(BASE/f'raw/duration_{s["duration_ticks"]}.csv');assert len(r)==5
            v=[float(x['control-fraction']) for x in r];h=2.7764451051977987*statistics.stdev(v)/math.sqrt(5)
            close(s['control_ci_low'],statistics.mean(v)-h);close(s['control_ci_high'],statistics.mean(v)+h)
    for name,seed in [('below_boundary',500003),('boundary',500121),('above_boundary',500207)]:
        trace=table(BASE/f'raw/representative_{name}.csv');assert len(trace)==1001
        match=next(r for r in rows if int(r['500000 + behaviorspace-run-number'])==seed)
        for m in ('control-fraction','controlled-kingdoms','mean-religion','mean-dependency'):close(trace[-1][m],match[m])
    print('PASS: hashed unchanged model/XML/raw; exact grid/seeds/settings; independent t/Wilson statistics; df=4 duration pilot; replay endpoints')
    return rows


def audit_artifacts():
    report=pymupdf.open(ROOT/'report/foundation_report.pdf');assert len(report)==2
    deck=pymupdf.open(ROOT/'presentation/foundation.pdf');assert len(deck)==6
    for doc in (report,deck):
        for page in doc:
            assert page.get_text().strip()
            for block in page.get_text('dict')['blocks']:
                if block['type']==0:
                    x0,y0,x1,y1=block['bbox'];assert x0>=0 and y0>=0 and x1<=page.rect.width+.5 and y1<=page.rect.height+.5
    prs=Presentation(ROOT/'presentation/foundation.pptx');assert len(prs.slides)==6
    for i,slide in enumerate(prs.slides):
        assert slide.notes_slide.notes_text_frame.text.strip()
        actual=Image.open(io.BytesIO(slide.shapes[0].image.blob)).convert('RGB')
        pix=deck[i].get_pixmap(matrix=pymupdf.Matrix(2,2))
        expected=Image.frombytes('RGB',(pix.width,pix.height),pix.samples)
        assert actual.size==expected.size and actual.tobytes()==expected.tobytes()
    assert slide.shapes[0].width==prs.slide_width
    gif=Image.open(MEDIA/'foundation_demo.gif');assert gif.n_frames==21
    assert prs.slides[2].shapes[1].image.ext=='gif'
    expected={'model/foundation.nlogo','experiments/tipping_confirmatory.xml','report/foundation_report.pdf','presentation/foundation.pptx','presentation/foundation.pdf'}
    with zipfile.ZipFile(ROOT/'dist/Foundation_Simulation_Submission.zip') as z:
        assert set(z.namelist())==expected and z.testzip() is None
        with tempfile.TemporaryDirectory(prefix='foundation-extracted-') as temp:
            z.extractall(temp)
            for name in expected:assert (Path(temp)/name).read_bytes()==(ROOT/name).read_bytes()
    print('PASS: 2 report pages; 6 slides in each format; identical slide rendering; real 21-frame GIF; five-file ZIP extracts byte-identically')


def runtime(home,rows):
    with tempfile.TemporaryDirectory(prefix='foundation-submission-runtime-') as t:
        work=Path(t)
        with zipfile.ZipFile(ROOT/'dist/Foundation_Simulation_Submission.zip') as z:z.extractall(work)
        # First cell/repetition: same seed and commands, all fixed parameters intact.
        tree=ET.parse(work/'experiments/tipping_confirmatory.xml');e=tree.getroot().find('experiment');e.set('repetitions','1')
        for s in e.findall('enumeratedValueSet'):
            for v in list(s)[1:]:s.remove(v)
        xml=work/'one-cell.xml';tree.write(xml,encoding='unicode')
        out=work/'probe.csv'
        cmd=[str(Path(home).resolve()/'NetLogo_Console'),'--headless','--model',str(work/'model/foundation.nlogo'),'--setup-file',str(xml),'--experiment',e.get('name'),'--threads','1','--table',str(out),'--update-plots']
        r=subprocess.run(cmd,capture_output=True,text=True,timeout=120)
        assert r.returncode==0 and 'RUNTIME ERROR' not in r.stdout+r.stderr,r.stdout+r.stderr
        observed=table(out);assert len(observed)==1
        expected=next(r for r in rows if r['[run number]']=='1')
        for m in (e.findtext('metric'), 'controlled-kingdoms','mean-religion','mean-dependency','cumulative-trade-profit','ticks'):close(observed[0][m],expected[m])
        assert observed[0]['model-valid?']=='true'
    print('PASS: extracted self-contained model + selected XML cell compile/run in NetLogo 6.4.0 with plots; seed 500001 matches recorded results')


MEDIA=ROOT/'presentation/media'
if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--netlogo-home');a=p.parse_args()
    rows=audit_science();audit_artifacts()
    if a.netlogo_home:runtime(a.netlogo_home,rows)
