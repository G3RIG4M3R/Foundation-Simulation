#!/usr/bin/env python3
"""Build the two-page report, six-slide PPTX/PDF and minimal submission ZIP.

Slides use one shared layout rendered to PDF and embedded in PPTX to preserve
classroom legibility across viewers. Speaker notes remain editable; slide 3
includes the genuine animated GUI GIF over its matching first-frame poster.
"""
from pathlib import Path
import tempfile
import zipfile

import pymupdf
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph
from reportlab.lib.styles import ParagraphStyle

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / 'results/tipping_point_investigation/figures'
MEDIA = ROOT / 'presentation/media'
NAVY, TEAL, GOLD = '#142b42', '#087f8c', '#bc7a20'
FONT = '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
pdfmetrics.registerFont(TTFont('D', FONT))
pdfmetrics.registerFont(TTFont('DB', FONT.replace('.ttf', '-Bold.ttf')))


def paragraph(c, text, x, top, width, size=9.5, color=NAVY, leading=None):
    style = ParagraphStyle('body', fontName='D', fontSize=size,
                           leading=leading or size*1.36, textColor=HexColor(color))
    p = Paragraph(text, style)
    _, height = p.wrap(width, 2000)
    if top-height < 30: raise ValueError('Content runs beyond bottom margin: '+text[:60])
    p.drawOn(c, x, top-height)
    return top-height-8


def image(c, path, x, top, width, height):
    c.drawImage(str(path), x, top-height, width=width, height=height,
                preserveAspectRatio=True, anchor='c', mask='auto')


def heading(c, text, x, top, size=13):
    c.setFont('DB', size); c.setFillColor(HexColor(TEAL)); c.drawString(x, top-size, text)
    return top-size-12


def report():
    path = ROOT/'report/foundation_report.pdf'
    c = canvas.Canvas(str(path), pagesize=(595.28,841.89))
    c.setTitle('Foundation: economic dependency and widespread control')
    x,w=43,509
    author=(ROOT/'report/author.txt').read_text().strip()
    y=heading(c,'Foundation: influence through economic dependency',x,800,15)
    y=paragraph(c,author+' • Collective Intelligence • Assignment 1 • Autumn 2026',x,y,w,8)
    y=heading(c,'Motivation and research question',x,y)
    y=paragraph(c,'Isaac Asimov’s <i>Foundation</i> inspires a fictional model of non-military expansion: missionaries maintain advanced devices and religious institutions, while traders create dependence on imported technology. How do trader population and trade attractiveness influence widespread Foundation control? The hypothesis is that more frequent, attractive trade raises dependency and trust enough to establish broad leverage under favorable missionary conditions. This is an original fictional abstraction, not a predictive representation of real societies.',x,y,w)
    y=heading(c,'Agents and interactions',x,y)
    y=paragraph(c,'The 64 × 64 toroidal world contains 31 stationary planets: Terminus, four kingdoms of five planets each (one capital per kingdom), and ten independent markets. Two traveling breeds interact with planets: missionaries and traders. Stochastic destination choice favors accessible, useful nearby worlds but includes exploration. Accepted missions increase religion with saturation; established temples support religion and device maintenance. Religion improves trade admission without making conversion mandatory. Successful sales raise technological dependency, infrastructure health and trust, create or reinforce trade routes, and generate revenue. Deaths and paid recruitment constrain staffing.',x,y,w)
    y=heading(c,'ODD overview and dynamic environment',x,y)
    y=paragraph(c,'<b>Purpose and state:</b> study emergent, reversible leverage; planets hold religion, dependency, demand, health, wealth, hostility and policy, travelers hold skill and destination, and routes hold strength and age. <b>Initialization:</b> a random separated galaxy and heterogeneous bounded planet states; open governments, no routes or controlled worlds; no external input data. <b>Scheduling:</b> environment → missions → trade → route wear → government/recruitment every ten ticks → local political effects → control → clock. <b>Design concepts:</b> saturation, stochastic admission, economic feedback, policy hysteresis and emergent networks. Demand regenerates, technology degrades, local substitutes reduce dependence and unused routes disappear. Capitals can impose restrictions or embargoes; dependency and falling wealth can induce concessions. Kingdom names do not script victory.',x,y,w)
    y=heading(c,'Control, parameters and interface',x,y)
    y=paragraph(c,'Planet leverage is clamp01(0.30 religion + 0.35 dependency + 0.25 trust + 0.15(1 − hostility)). Leverage ≥0.58 for five consecutive ticks establishes effective control; a lower tick removes it. A kingdom requires its controlled capital and ≥60% of its population on controlled worlds. Control is influence, not sovereignty. Control fraction counts the 30 external planets only.',x,y,w)
    y=paragraph(c,'Key controls are missionary effectiveness, missionary staffing, trader staffing, trade attractiveness, royal intolerance, technology decay, religion-assisted trade admission and local substitution. The confirmatory settings are effectiveness 0.50, missionaries 12, intolerance 0, decay 0.015, religion-trade weight 0.80 and independence effort 0.15; traders and attractiveness vary on page 2. The ninth slider sets duration: interactive default 0 means unlimited, while experiments capture a finite limit at setup.',x,y,w)
    y=paragraph(c,'The self-contained NetLogo 6.4 model provides setup, go-once and go; four planet views; twelve monitors; control, influence and trade time series; and a religion histogram. The Info tab supplies the full ODD and equations. No warfare, empirical calibration or individual worshippers are modeled; space and credits are abstract.',x,y,w)
    y=paragraph(c,'<b>References and authorship:</b> Asimov, I. (1951), <i>Foundation</i> (literary inspiration); Wilensky, U. (1999), <i>NetLogo</i>, Northwestern University; Takács, T. (2026), Assignment 1. AI assistance was used for code, documentation and delivery.',x,y,w,7.8)
    # Conceptual mechanism diagram, not a simulated measurement.
    labels=['Missions → religion', 'Trade → dependency', 'Leverage ↔ resistance']
    for i,label in enumerate(labels):
        bx=x+i*173
        c.setFillColor(HexColor('#edf5f6'));c.rect(bx,68,163,30,fill=1,stroke=0)
        c.setFillColor(HexColor(TEAL));c.setFont('D',8);c.drawCentredString(bx+81.5,80,label)
    c.setFont('D',8);c.drawRightString(552,25,'1 / 2');c.showPage()
    y=heading(c,'Experiment, measured results and interpretation',x,800,15)
    y=paragraph(c,'BehaviorSpace varies <b>initial-traders</b> {8, 12, 16, 20, 24} × <b>trade-attractiveness</b> {0.60, 0.80, 1.00}: 15 cells, 20 repetitions each, 300 runs at 1,000 ticks. XML: <b>tipping_confirmatory.xml</b>; seeds 500000 + run number. Fixed settings are on page 1. The earlier 980-run missionary × intolerance sweep found no widespread control; separate duration, mechanism-screening and boundary-search experiments selected these ranges. Exploratory and confirmatory observations are not pooled.',x,y,w,8.7)
    y=paragraph(c,'Primary outcomes are final control fraction and controlled kingdoms (0–4); per-cell means and 95% Student t intervals use df=19. Widespread control means final fraction ≥0.50 (≥15/30 worlds); its probability uses 95% Wilson binomial intervals.',x,y,w,8.7)
    image(c,FIG/'control_heatmap.png',x+30,y,449,240);y-=242
    y=paragraph(c,'Figure 1. Mean final control fraction; each cell shows mean ±95% t CI half-width. Color spans 0–1. The 50% criterion is an operational definition, not a universal constant.',x,y,w,7.7)
    image(c,FIG/'representative_control_timeseries.png',x,y,w,153);y-=155
    y=paragraph(c,'Figure 2. Genuine below/near/above-boundary confirmatory replays. Dashed line: 50% criterion. Seeds 500003, 500121, 500207; final states 0.200, 0.500, 0.633. The last case uses attractiveness 0.80; the other two use 0.60.',x,y,w,7.7)
    y=paragraph(c,'At attractiveness <b>0.60</b>, 12 traders yield <b>0/20</b> widespread runs and mean 0.387 [0.364, 0.409]; 16 yield <b>10/20</b> and 0.497 [0.460, 0.533]; 20 yield <b>20/20</b> and 0.605 [0.577, 0.633]. Corresponding Wilson intervals are [0.000, 0.161], [0.299, 0.701], [0.839, 1.000]. At 0.80, 12 traders yield 16/20, moving the boundary toward fewer traders. More trade increases dependency and trust, strengthening leverage. Kingdom-count means and intervals are retained in the processed CSV.',x,y,w,8.5)
    y=paragraph(c,'The crossing probability changes rapidly between 12 and 20 traders at 0.60, while continuous control grows more smoothly: an <b>operational stochastic threshold</b>, not proof of a mathematical phase transition. The near-boundary replay first crosses at tick 967 and is widespread in only 26/1,001 states; terminal crossing does not establish persistence. The above-boundary replay is widespread throughout the final 100 ticks.',x,y,w,8.5)
    y=paragraph(c,'<b>Takeaway:</b> favorable missionary access enables expansion, but sufficiently strong economic ties drive widespread control in this model. <b>Limits:</b> fictional coefficients, favorable fixed politics, finite horizon, coarse four-trader spacing, twenty runs and selection after exploratory search restrict generality. Future work: fresh 14/18-trader validation cells and persistence measurements across all runs.',x,y,w,8.5)
    c.setFont('D',8);c.drawRightString(552,25,'2 / 2');c.save()


NOTES = [
('0:00–0:50', 'Asimov imagined a small technologically advanced Foundation gaining influence without military conquest. We abstract this into religion and commerce, rather than reenacting the novels. Ask the question on the slide. Our hypothesis is that frequent, attractive trade creates dependency under favorable access. The result describes this fictional simulation, not real societies.'),
('0:50–1:45', 'Planets are stationary; missionaries and traders travel. Missions saturate religious influence, and temples maintain technology. Sales build dependency and trust, while links connect markets. Technology wears out and demand returns. Capitals react with restrictions, but economic pressure can soften policies. Control is reversible leverage: five ticks above the local criterion, and a controlled capital plus 60 percent population for a kingdom. Point out the genuine interface, duration slider, plots and histogram.'),
('1:45–2:40', 'The animation and stills come from the actual NetLogo GUI with seed 500207, 20 traders and attractiveness 0.80. GIF time is accelerated. Show initial, middle and terminal states; describe the color legend. Religion opens access, then economic ties and control accumulate. This replay crosses 50 percent at tick 620 and finishes with 19 of 30 planets and two kingdoms. Replay is illustrative, not an extra independent observation.'),
('2:40–3:30', 'The original 49-cell missionary versus intolerance experiment found no widespread control in 980 runs at 450 ticks. Duration tests, factorial screening and a boundary search pointed toward economic variables. We then used a separate 15-cell design with 20 unique seeds per cell and a finite 1000-tick horizon. Settings are favorable and fixed. Explain Student t intervals for continuous means and Wilson intervals for crossing probability. We do not pool search data with confirmation.'),
('3:30–4:35', 'Read the bottom heatmap row at attractiveness 0.60. Twelve traders: 0 of 20 widespread runs, mean control 0.387. Sixteen: 10 of 20, mean 0.497. Twenty: 20 of 20, mean 0.605. Zero and twenty successes still have uncertainty: Wilson upper 0.161 and lower 0.839. At 0.80, twelve traders produce 16 of 20 widespread runs. Mean control grows more smoothly than the threshold-crossing probability. A 50 percent cutoff makes the transition appear sharper.'),
('4:35–5:30', 'Dependency and trust are major mechanisms, under favorable missionary and political settings. This is a stochastic operational threshold, not proof of a mathematical phase transition. The boundary replay crosses late and fluctuates: only 26 of 1001 observed states are widespread. Terminal outcomes do not prove lasting control. End with the takeaway. Future work should use fresh validation seeds at 14 and 18 traders and record persistence for every run. Invite questions.')]


def slides():
    path=ROOT/'presentation/foundation.pdf'
    c=canvas.Canvas(str(path),pagesize=(960,540));c.setTitle('Foundation — economic dependency')
    def base(n,title,kicker):
        c.setFillColor(HexColor('#f4f7fa'));c.rect(0,0,960,540,fill=1,stroke=0)
        c.setFillColor(HexColor(NAVY));c.rect(0,453,960,87,fill=1,stroke=0)
        paragraph(c,title,36,518,890,27,'#ffffff',32)
        paragraph(c,kicker,38,472,880,11,'#a8dbe0')
        c.setFillColor(HexColor(TEAL));c.rect(36,27,888,2,fill=1,stroke=0)
        c.setFont('D',8);c.setFillColor(HexColor(NAVY));c.drawString(36,12,'Collective Intelligence • Autumn 2026     |     Foundation')
        c.setFont('D',9);c.drawRightString(922,12,str(n)+' / 6')
    base(1,'Can commerce create widespread control?','A fictional agent-based model inspired by Isaac Asimov’s Foundation')
    paragraph(c,'A small technological center<br/>Two paths to influence<br/><b>Religion + economic dependency</b>',42,410,460,28,leading=40)
    paragraph(c,'How do trader population and trade attractiveness influence the emergence of widespread Foundation control?',42,242,500,23,leading=31)
    image(c,MEDIA/'view-1000.png',590,415,310,310)
    paragraph(c,'Hypothesis: more frequent, attractive trade builds dependency and trust under favorable access.',42,121,520,17)
    paragraph(c,(ROOT/'report/author.txt').read_text().strip(),590,89,320,11)
    c.showPage()
    base(2,'Three breeds, interacting mechanisms','Original rules • 31 planets • 64 × 64 wrapping world • reversible influence')
    image(c,MEDIA/'interface-1000.png',35,420,560,355)
    paragraph(c,'<b>Missionaries</b><br/>Conversion saturates; temples support maintenance.<br/><br/><b>Traders</b><br/>Sales build dependency, trust and routes.<br/><br/><b>Planets / capitals</b><br/>Demand regenerates; technology wears; governments resist or concede.',625,420,300,16,leading=22)
    paragraph(c,'Control: sustained leverage ≥0.58.<br/>Kingdom: capital + ≥60% population.<br/>Twelve monitors + time series + histogram.',625,134,300,13)
    c.showPage()
    base(3,'Watch influence develop','Genuine NetLogo GUI replay • seed 500207 • 20 traders • attractiveness 0.80')
    for x,tick in [(38,0),(338,500),(638,1000)]:
        image(c,MEDIA/f'view-{tick:04d}.png',x,420,280,280)
        label = 'Replay: ticks 0–1,000' if tick == 0 else f'Tick {tick}'
        paragraph(c,f'<b>{label}</b>',x,130,280,17)
    paragraph(c,'Control view: gold = Foundation; green = controlled; gray = uncontrolled. Links show trade routes.',38,91,884,15)
    paragraph(c,'First ≥50% at tick 620 → 19/30 controlled worlds and 2 kingdoms at tick 1,000. Animation is accelerated.',38,62,884,14)
    c.showPage()
    base(4,'A separate confirmatory two-factor experiment','Exploratory evidence selected economic factors; historical results remain separate')
    paragraph(c,'<b>5 trader populations</b><br/>8 · 12 · 16 · 20 · 24<br/><br/><b>3 trade attractiveness levels</b><br/>0.60 · 0.80 · 1.00',42,409,430,26)
    paragraph(c,'<b>15 cells × 20 repetitions = 300 runs</b><br/>1,000 ticks per run<br/>Unique seeds 500001–500300<br/><br/>Outcomes: final control fraction and controlled kingdoms.',505,409,413,22)
    paragraph(c,'Fixed: effectiveness 0.50; intolerance 0; missionaries 12; decay 0.015; religion-trade weight 0.80; substitution 0.15.',42,197,876,17)
    paragraph(c,'Mean uncertainty: 95% Student t (df=19).<br/>Widespread probability: 95% Wilson interval; final control ≥50%.',42,139,876,17)
    paragraph(c,'Earlier 980-run sweep: no widespread control. Duration pilot → mechanism screen → boundary search → confirmation.',42,77,876,15)
    c.showPage()
    base(5,'A rapid change in crossing probability','Continuous mean control and binary widespread probability are distinct outcomes')
    image(c,FIG/'control_heatmap.png',32,426,548,343)
    paragraph(c,'At attractiveness <b>0.60</b><br/><br/><b>12 traders: 0/20</b><br/>Mean control 0.387<br/><br/><b>16 traders: 10/20</b><br/>Mean control 0.497<br/><br/><b>20 traders: 20/20</b><br/>Mean control 0.605',611,422,310,19,leading=24)
    paragraph(c,'Heatmap: mean ±95% t CI half-width.<br/>Wilson intervals: 0/20 [0, .161]; 10/20 [.299, .701]; 20/20 [.839, 1].',40,93,870,12)
    paragraph(c,'At attractiveness 0.80, 12 traders reach widespread control in 16/20 runs: the apparent boundary shifts left.',40,52,875,12)
    c.showPage()
    base(6,'Economic ties drive expansion under favorable access','Operational stochastic threshold • finite-horizon evidence • fictional abstraction')
    paragraph(c,'<b>Mechanism</b><br/>Trade → dependency + trust → sustained leverage.<br/>Missionary influence alone was insufficient in the original tested settings.',40,415,440,21)
    paragraph(c,'<b>What the threshold means</b><br/>A rapid probability change around a defined 50% criterion.<br/>Continuous control rises more smoothly; no proven mathematical phase transition.',510,415,408,21)
    paragraph(c,'<b>Limits and next tests</b><br/>Favorable politics and religion; coarse grid; twenty runs.<br/>Near-boundary crossing can be transient (26/1,001 states in the replay).<br/>Test 14/18 traders with fresh seeds; measure persistence in all runs.',40,235,874,19)
    paragraph(c,'Takeaway: strong economic ties can produce widespread Foundation control in this model.',40,111,874,18,TEAL)
    paragraph(c,'Sources: Asimov (1951); Wilensky (1999); Takács (2026). AI-assisted development.',40,50,874,8)
    c.showPage();c.save()
    prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
    with tempfile.TemporaryDirectory(prefix='foundation-slides-') as tmp:
        pdf=pymupdf.open(path)
        for i,page in enumerate(pdf):
            png=Path(tmp)/f'slide-{i}.png'
            page.get_pixmap(matrix=pymupdf.Matrix(2,2)).save(png)
            slide=prs.slides.add_slide(prs.slide_layouts[6])
            slide.shapes.add_picture(str(png),0,0,width=prs.slide_width,height=prs.slide_height)
            if i==2:
                slide.shapes.add_picture(str(MEDIA/'foundation_demo.gif'),
                                        Inches(38/72), Inches(120/72),
                                        width=Inches(280/72),height=Inches(280/72))
            slide.notes_slide.notes_text_frame.text=NOTES[i][0]+'\n'+NOTES[i][1]
        prs.save(ROOT/'presentation/foundation.pptx')
    (ROOT/'presentation/rehearsal.md').write_text('# Rehearsal script (target 5:30)\n\n'+
        '\n\n'.join(f'## Slide {i+1} — {timing}\n\n{text}' for i,(timing,text) in enumerate(NOTES))+
        '\n\nThe PDF shows stills; the PPTX includes the genuine GIF over the tick-0 poster.\nTest animation in PowerPoint slideshow mode. If a viewer does not animate GIFs,\nopen `media/foundation_demo.gif` separately; the visible sequence still supports the talk.\nSlide text is rendered from the same PDF layout for consistent appearance.\nEdit content in `scripts/build_delivery.py`, then rebuild.\n')


def package():
    files=['model/foundation.nlogo','experiments/tipping_confirmatory.xml',
           'report/foundation_report.pdf','presentation/foundation.pptx','presentation/foundation.pdf']
    dist=ROOT/'dist';dist.mkdir(exist_ok=True)
    with zipfile.ZipFile(dist/'Foundation_Simulation_Submission.zip','w',zipfile.ZIP_DEFLATED) as z:
        for path in files: z.write(ROOT/path,path)
    print('Built two-page report, six-slide PDF/PPTX and five-file ZIP.')


if __name__=='__main__':
    report();slides();package()
